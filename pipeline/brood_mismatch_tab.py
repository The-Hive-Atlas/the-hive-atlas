"""Add a "Brood mismatches" tab to a cleaned antibiotic workbook: antibiotic tests whose brood submission
(same Sample# and year) is recorded entirely AFB-negative, with a match check on state and month.

Usage: python3 pipeline/brood_mismatch_tab.py Antibiotics_cleaned.xlsx Brood_cleaned.xlsx
Then open and save the workbook in Excel or LibreOffice to recalculate the TE and TY formulas.
"""
import sys, pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
A, B = sys.argv[1], sys.argv[2]
a=pd.read_excel(A,sheet_name=0); a['row']=a.index+2
b=pd.read_excel(B); b['row']=b.index+2
b['sn']=b['Sample#'].astype(str); a['sn']=a['Sample#'].astype(str).str.split('-').str[0]
M={m:i for i,m in enumerate(['','Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'])}
def rng(rows):
    rows=sorted(rows); out=[]; s=p=rows[0]
    for r in rows[1:]:
        if r==p+1: p=r; continue
        out.append(f"{s}-{p}" if s!=p else str(s)); s=p=r
    out.append(f"{s}-{p}" if s!=p else str(s)); return ', '.join(out)
bg={k:g for k,g in b.groupby(['sn','Year'])}; ntests=a.groupby(['sn','Year']).size()
recs=[]
for _,r in a.iterrows():
    g=bg.get((r.sn,r.Year))
    if g is None or g.AFB.max()!=0: continue
    f=g.iloc[0]; day='' if pd.isna(f.Day) else int(f.Day)
    sa=r.State if isinstance(r.State,str) else 'na'
    probs=[]
    if sa!=f.State: probs.append('state differs')
    if r.DateDiagnosed.month!=M.get(str(f.Month)[:3]): probs.append('month differs')
    chk='Strong: same state and month' if not probs else 'Weak: '+' and '.join(probs)+' (may be a different sample)'
    recs.append([r.row,r['Sample#'],r.sn,r.DateDiagnosed.date(),r.Year,sa,r.TEmm,r.TYmm,int(ntests[(r.sn,r.Year)]),len(g),int(g.AFB.sum()),int(g.EFB.sum()),f"{f.Month or ''} {day} {int(f.Year)}".strip(),f.State,rng(g.row.tolist()),chk])
recs.sort(key=lambda x:(x[4],x[2],str(x[1])))
hdr=['Antibiotic Excel row','Antibiotic Sample#','Submission (base Sample#)','DateDiagnosed','Year','State (antibiotic)','TEmm','TYmm','Antibiotic tests in submission','Brood rows in submission','Brood AFB positives','Brood EFB positives','Brood collection date','State (brood)','Brood Excel rows (Brood_Oct28-2026.xlsx)','Match check']
wb=openpyxl.load_workbook(A)
if 'Brood mismatches' in wb.sheetnames: del wb['Brood mismatches']
ws=wb.create_sheet('Brood mismatches',1)
ws.append(['Antibiotic tests whose matching brood submission (same Sample# and year) has every colony recorded AFB = 0. Antibiotic testing is only done on AFB-positive samples, so these brood records likely miss positives. Brood row numbers refer to the original Brood_Oct28-2026.xlsx (same rows in the cleaned file). Rows shaded pink are weak matches: the state or month differs, so they may be a different sample that happens to share the number.'])
ws.append(hdr)
pink=PatternFill('solid',fgColor='F7E1DB')
for x in recs:
    ws.append(x)
    if x[-1].startswith('Weak'):
        for c in ws[ws.max_row]: c.fill=pink
ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(hdr))
ws['A1'].font=Font(italic=True); ws['A1'].alignment=Alignment(wrap_text=True,vertical='top'); ws.row_dimensions[1].height=46
for c in ws[2]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='F3E3BF'); c.alignment=Alignment(wrap_text=True,vertical='bottom')
ws.row_dimensions[2].height=45; ws.freeze_panes='A3'; ws.auto_filter.ref=f"A2:{L(len(hdr))}{ws.max_row}"
for i,w in enumerate([11,14,13,13,7,10,8,8,12,11,11,11,15,10,34,30],1): ws.column_dimensions[L(i)].width=w
wb.save(A)
print(len(recs),'rows; strong',sum(x[-1].startswith('Strong') for x in recs),'weak',sum(x[-1].startswith('Weak') for x in recs))
