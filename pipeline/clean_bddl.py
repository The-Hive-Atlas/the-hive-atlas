"""Clean a BDDL adult-bee spreadsheet before building the map.

Usage: python3 pipeline/clean_bddl.py Adult_BDDL_1984Present.xlsx Adult_BDDL_cleaned.xlsx

Rules (agreed October 2026):
  VIL above 0 but VD 0       -> VD = 1
  VD 1 but VIL 0             -> VIL = 1.0 (placeholder: positive but none counted in the wash)
  MSPB above 0 but NOS 0     -> NOS = 1
  NOS 1 but MSPB 0           -> MSPB = 0.1 (placeholder: positive but below the counting range)
  State codes trimmed and upper-cased; KT -> KY. Extra spaces trimmed from zip codes.
"na" (not counted) is never changed. Adds a "Cleaning log" tab (every change, with the original
value) and a "For review" tab (positives with no count, and zip codes that point to another state).
Needs: pip install openpyxl zipcodes
"""
import openpyxl, re, sys, zipcodes
from openpyxl.styles import Font
src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src); ws = wb['Adult']
h = [c.value for c in ws[1]]; C = {k: h.index(k) + 1 for k in h}
log, review = [], []
def isnum(v): return isinstance(v, (int, float)) and not isinstance(v, bool)
def setv(r, col, new, rule):
    cell = ws.cell(r, C[col]); log.append((r, ws.cell(r, C['Sample#']).value, ws.cell(r, C['Year']).value, col, cell.value, new, rule)); cell.value = new
for r in range(2, ws.max_row + 1):
    g = lambda k: ws.cell(r, C[k]).value
    vd, vil, nos, ms = g('VD'), g('VIL'), g('NOS'), g('MSPB')
    if vd == 0 and isnum(vil) and vil > 0: setv(r, 'VD', 1, 'VIL above 0 but VD was 0: set VD to 1')
    if vd == 1 and isnum(vil) and vil == 0: setv(r, 'VIL', 1.0, 'VD was 1 but VIL was 0: set VIL to 1.0')
    if nos == 0 and isnum(ms) and ms > 0: setv(r, 'NOS', 1, 'MSPB above 0 but NOS was 0: set NOS to 1')
    if nos == 1 and isnum(ms) and ms == 0: setv(r, 'MSPB', 0.1, 'NOS was 1 but MSPB was 0: set MSPB to 0.1')
    st = g('State')
    if isinstance(st, str):
        u = {'KT': 'KY'}.get(st.strip().upper(), st.strip().upper())
        new = u if len(u) == 2 and u.isalpha() and u != 'NA' else st.strip()  # leave 'na' and odd codes as they are
        if new != st: setv(r, 'State', new, 'State code fixed (extra space or nonstandard code)')
    z = g('ZipCode')
    if isinstance(z, str) and z != z.strip() and z.strip():
        setv(r, 'ZipCode', z.strip(), 'Extra spaces removed from zip code')
    # Review items (not changed)
    if g('Year') >= 2015 and vd == 1 and vil == 'na': review.append((r, g('Sample#'), g('Year'), g('State'), 'VD is 1 but VIL is "na" (positive, no count recorded)', ''))
    if g('Year') >= 2015 and nos == 1 and ms == 'na': review.append((r, g('Sample#'), g('Year'), g('State'), 'NOS is 1 but MSPB is "na" (positive, no count recorded)', ''))
    z = g('ZipCode'); st = g('State')
    if z is not None:
        m = re.match(r'\s*(\d{3,5})', str(z))
        if m:
            raw = m.group(1); hit = zipcodes.matching(raw.zfill(5))
            if hit and hit[0]['state'] != st:
                review.append((r, g('Sample#'), g('Year'), st, 'Zip code points to a different state', '%s is %s, %s%s' % (z, hit[0]['city'], hit[0]['state'], ' (zip is missing digits)' if len(raw) < 5 else '')))
ws2 = wb.create_sheet('Cleaning log')
ws2.append(['Excel row', 'Sample#', 'Year', 'Column', 'Original value', 'New value', 'Rule'])
for x in log: ws2.append(list(x))
ws3 = wb.create_sheet('For review')
ws3.append(['Excel row', 'Sample#', 'Year', 'State', 'Issue (not changed)', 'Detail'])
for x in review: ws3.append(list(x))
for s, w in ((ws2, [10, 10, 7, 10, 14, 11, 52]), (ws3, [10, 10, 7, 7, 52, 44])):
    for c in s[1]: c.font = Font(bold=True)
    s.freeze_panes = 'A2'
    for i, wd in enumerate(w): s.column_dimensions[chr(65 + i)].width = wd
wb.save(dst)
from collections import Counter
print(Counter(x[6] for x in log)); print(Counter(x[4] for x in review))
