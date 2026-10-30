"""Clean a BDDL antibiotic-resistance spreadsheet (American foulbrood, tetracycline and tylosin).

Usage: python3 pipeline/clean_antibiotics.py Antibiotics.xlsx Antibiotics_cleaned.xlsx [Brood_cleaned.xlsx]
Needs: pip install openpyxl. Recalculate the TE and TY formulas afterwards (open and save in Excel or LibreOffice).

Each row is one AFB-positive sample whose Paenibacillus larvae culture was tested by agar disc
diffusion. TEmm and TYmm are inhibition-zone diameters in mm (0 to 84). Resistance (%) = 100 - 50x/42.

Changes made (each listed on a "Cleaning log" tab with the original value):
  - The second "TEmm" header is renamed "TYmm" (it holds tylosin zones; the TY formula reads it).
  - Empty, unnamed columns are removed.
  - The empty Month column is filled from DateDiagnosed (diagnosis and collection fall in the same month).
  - State codes are trimmed and upper-cased; PEI -> PE and ALB -> AB (Canadian provinces).
  - With a brood file: a state coded "N" is filled from the matching brood record.

Flagged but not changed ("For review" tab): unknown state codes, missing sample numbers, exact
duplicate rows, zone sizes of 0 or below the 6 mm disc, outlying zones, Canadian samples, and
(with a brood file) antibiotic tests whose matching brood record says AFB was not found.
"""
import calendar
import sys
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font

US = set('AL AK AZ AR CA CO CT DE DC FL GA HI ID IL IN IA KS KY LA ME MD MA MI MN MS MO MT NE NV NH NJ NM NY NC ND OH OK OR PA RI SC SD TN TX UT VT VA WA WV WI WY PR'.split())
CANADA = set('AB BC MB NB NL NS ON PE QC SK'.split())
FIX = {'PEI': 'PE', 'ALB': 'AB'}


def main(src, dst, brood=None):
    wb = openpyxl.load_workbook(src)
    ws = wb.worksheets[0]
    log, review = [], []
    # header fix: second TEmm is tylosin
    heads = [c.value for c in ws[1]]
    idx = [i for i, h in enumerate(heads) if h == 'TEmm']
    if len(idx) == 2:
        ws.cell(1, idx[1] + 1).value = 'TYmm'
        log.append(('header', '', '', 'column %d' % (idx[1] + 1), 'TEmm', 'TYmm', 'Second TEmm column holds tylosin zones; renamed TYmm'))
    for c in range(ws.max_column, 0, -1):
        if ws.cell(1, c).value is None and all(ws.cell(r, c).value is None for r in range(2, ws.max_row + 1)):
            ws.delete_cols(c)
            log.append(('all', '', '', 'column %d' % c, '(empty, no header)', '(removed)', 'Empty unnamed column removed'))
    h = [c.value for c in ws[1]]
    C = {k: h.index(k) + 1 for k in h if k}
    # Deleting columns does not update formulas, so rewrite TE and TY to point at the zone columns
    from openpyxl.utils import get_column_letter as L
    for r in range(2, ws.max_row + 1):
        for res, mm in (('TE', 'TEmm'), ('TY', 'TYmm')):
            v = ws.cell(r, C[mm]).value
            ws.cell(r, C[res]).value = '=100-(50*%s%d)/42' % (L(C[mm]), r) if isinstance(v, (int, float)) else 'na'

    bstate, bafb = {}, {}
    if brood:
        bws = openpyxl.load_workbook(brood, read_only=True).worksheets[0]
        it = bws.iter_rows(values_only=True)
        bh = list(next(it))
        bi = {k: bh.index(k) for k in ('Sample#', 'Year', 'State', 'AFB')}
        for r in it:
            k = (str(r[bi['Sample#']]), r[bi['Year']])
            bstate.setdefault(k, r[bi['State']])
            bafb[k] = max(bafb.get(k, 0), r[bi['AFB']] or 0)

    def setv(r, col, new, rule):
        cell = ws.cell(r, C[col])
        log.append((r, ws.cell(r, C['Sample#']).value, ws.cell(r, C['Year']).value, col, repr(cell.value), new, rule))
        cell.value = new

    def flag(r, issue, detail=''):
        g = lambda k: ws.cell(r, C[k]).value
        review.append((r, g('Sample#'), g('Year'), g('State'), issue, detail))

    seen = defaultdict(list)
    for r in range(2, ws.max_row + 1):
        g = lambda k: ws.cell(r, C[k]).value
        dd = g('DateDiagnosed')
        if g('Month') is None and hasattr(dd, 'month'):
            setv(r, 'Month', calendar.month_abbr[dd.month], 'Month filled from DateDiagnosed')
        st = g('State')
        if isinstance(st, str):
            u = FIX.get(st.strip().upper(), st.strip().upper())
            # Only standardize real state or province codes; leave 'na' (not recorded) and unknown codes as they are
            new = u if (u in US or u in CANADA) else st.strip()
            if new != st:
                setv(r, 'State', new, 'State code fixed (spacing, case, or province abbreviation)')
        st = g('State')
        base = str(g('Sample#')).split('-')[0]
        if st == 'N' and (base, g('Year')) in bstate:
            setv(r, 'State', bstate[(base, g('Year'))], 'State "N" filled from the matching brood record')
            st = g('State')
        if g('Sample#') is None:
            flag(r, 'Missing sample number')
        if st in CANADA:
            flag(r, 'Canadian sample (not shown on the U.S. map)', st)
        elif st not in US:
            flag(r, 'Unknown state code', repr(st))
        te, ty = g('TEmm'), g('TYmm')
        if isinstance(te, (int, float)) and te < 6:
            flag(r, 'Tetracycline zone of %s mm (0 = no inhibition; under 6 mm is smaller than the disc)' % te)
        if isinstance(ty, (int, float)) and ty < 30:
            flag(r, 'Tylosin zone of %s mm is far below all other tylosin zones (35 mm or more)' % ty)
        seen[(str(g('Sample#')), str(dd), st, te, ty)].append(r)
        if brood and bafb.get((base, g('Year'))) == 0:
            flag(r, 'Antibiotic test exists, but the matching brood record says AFB was not found', 'Sample# %s, %s' % (base, g('Year')))
    for k, rows in seen.items():
        if len(rows) > 1:
            for r in rows:
                flag(r, 'Exact duplicate of another row', 'rows ' + ', '.join(map(str, rows)))

    ws2 = wb.create_sheet('Cleaning log')
    ws2.append(['Excel row', 'Sample#', 'Year', 'Column', 'Original value', 'New value', 'Rule'])
    for x in log:
        ws2.append(list(x))
    ws3 = wb.create_sheet('For review')
    ws3.append(['Excel row', 'Sample#', 'Year', 'State', 'Issue (not changed)', 'Detail'])
    for x in review:
        ws3.append(list(x))
    for s, w in ((ws2, [10, 12, 7, 10, 16, 10, 58]), (ws3, [10, 12, 7, 7, 70, 30])):
        for c in s[1]:
            c.font = Font(bold=True)
        s.freeze_panes = 'A2'
        for i, wd in enumerate(w):
            s.column_dimensions[chr(65 + i)].width = wd
    wb.save(dst)
    from collections import Counter
    print(Counter(x[6] for x in log))
    print(Counter(x[4].split(' (')[0] if x[4].startswith('Canadian') else x[4][:60] for x in review))


if __name__ == '__main__':
    main(*sys.argv[1:])
