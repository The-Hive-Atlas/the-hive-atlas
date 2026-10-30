"""Clean a BDDL brood-sample spreadsheet (American and European foulbrood) before building the map.

Usage: python3 pipeline/clean_brood.py Brood.xlsx Brood_cleaned.xlsx
Needs: pip install openpyxl zipcodes

Changes made (each one listed on a "Cleaning log" tab with the original value):
  - Empty, unnamed columns are removed.
  - State codes are trimmed of spaces and tabs and upper-cased; KT -> KY.
  - Extra spaces are trimmed from zip codes.
  - A Day of 0 (not a real date) is cleared; the month and year are kept.

Flagged but not changed (listed on a "For review" tab):
  - Collection month later than the DateSubmitted date.
  - Zip codes that point to a different state than the State column.
AFB and EFB results are never changed.
"""
import calendar
import re
import sys

import openpyxl
import zipcodes
from openpyxl.styles import Font

MONTHS = {m: i for i, m in enumerate(calendar.month_abbr) if m}


def main(src, dst):
    wb = openpyxl.load_workbook(src)
    ws = wb.worksheets[0]
    log, review = [], []
    # remove empty, unnamed columns (right to left so indexes stay valid)
    for c in range(ws.max_column, 0, -1):
        if ws.cell(1, c).value is None and all(ws.cell(r, c).value is None for r in range(2, ws.max_row + 1)):
            ws.delete_cols(c)
            log.append(('all', '', '', 'column %d' % c, '(empty, no header)', '(removed)', 'Empty unnamed column removed'))
    h = [c.value for c in ws[1]]
    C = {k: h.index(k) + 1 for k in h if k}

    def setv(r, col, new, rule):
        cell = ws.cell(r, C[col])
        log.append((r, ws.cell(r, C['Sample#']).value, ws.cell(r, C['Year']).value, col, repr(cell.value), new, rule))
        cell.value = new

    for r in range(2, ws.max_row + 1):
        g = lambda k: ws.cell(r, C[k]).value
        st = g('State')
        if isinstance(st, str):
            u = {'KT': 'KY'}.get(st.strip().upper(), st.strip().upper())
            new = u if len(u) == 2 and u.isalpha() and u != 'NA' else st.strip()  # leave 'na' and odd codes as they are
            if new != st:
                setv(r, 'State', new, 'State code fixed (extra space, tab, or nonstandard code)')
        z = g('ZipCode')
        if isinstance(z, str) and z != z.strip() and z.strip():
            setv(r, 'ZipCode', z.strip(), 'Extra spaces removed from zip code')
        if g('Day') == 0:
            setv(r, 'Day', None, 'Day 0 is not a real date; cleared (month and year kept)')
        # review: collection month after the submission date
        ds, mo = g('DateSubmitted'), g('Month')
        if ds is not None and hasattr(ds, 'month') and isinstance(mo, str) and MONTHS.get(mo.strip()[:3].title()):
            m = MONTHS[mo.strip()[:3].title()]
            if ds.year == g('Year') and m > ds.month:
                review.append((r, g('Sample#'), g('Year'), g('State'), 'Collection month is later than DateSubmitted', '%s %s vs submitted %s' % (mo, g('Year'), ds.date())))
        z = g('ZipCode')
        if z is not None:
            m5 = re.match(r'\s*(\d{3,5})', str(z))
            if m5:
                raw = m5.group(1)
                hit = zipcodes.matching(raw.zfill(5))
                if hit and hit[0]['state'] != g('State'):
                    review.append((r, g('Sample#'), g('Year'), g('State'), 'Zip code points to a different state',
                                   '%s is %s, %s%s' % (z, hit[0]['city'], hit[0]['state'], ' (zip is missing digits)' if len(raw) < 5 else '')))

    ws2 = wb.create_sheet('Cleaning log')
    ws2.append(['Excel row', 'Sample#', 'Year', 'Column', 'Original value', 'New value', 'Rule'])
    for x in log:
        ws2.append(list(x))
    ws3 = wb.create_sheet('For review')
    ws3.append(['Excel row', 'Sample#', 'Year', 'State', 'Issue (not changed)', 'Detail'])
    for x in review:
        ws3.append(list(x))
    for s, w in ((ws2, [10, 10, 7, 10, 18, 12, 55]), (ws3, [10, 10, 7, 7, 46, 44])):
        for c in s[1]:
            c.font = Font(bold=True)
        s.freeze_panes = 'A2'
        for i, wd in enumerate(w):
            s.column_dimensions[chr(65 + i)].width = wd
    wb.save(dst)
    from collections import Counter
    print(Counter(x[6] for x in log))
    print(Counter(x[4] for x in review))


if __name__ == '__main__':
    main(*sys.argv[1:])
